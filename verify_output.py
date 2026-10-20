# -*- coding: utf-8 -*-
"""
Verification Suite for Serkom Laravel Guide Outputs
Validates HTML, DOCX, and PDF outputs for:
1. File existence & size
2. Dark Theme (#1e1e1e & #18181b)
3. Zero lexer token errors
4. High-contrast syntax highlighting & Blade in-tag expressions
5. Complete, untruncated code with clear modification markers
6. Proper DOCX cell widths (6.77in) and run-level color diversity (>= 12 colors)
7. Clean PDF pagination (33 pages, no single-line orphan page)
"""

import os
import re
import docx
from docx.shared import Inches
import PyPDF2
from pygments.token import Token

import build_guide
import content_data

def test_verification():
    print("=== TESTING COMPREHENSIVE VERIFICATION SUITE ===")
    
    # 1. Check generated files existence and sizes
    expected_files = [
        'PANDUAN_SERKOM_LARAVEL_SUPER_SIMPEL.docx',
        'PANDUAN_SERKOM_LARAVEL_SUPER_SIMPEL.pdf',
        'PANDUAN_UJIAN_CRUD_LARAVEL_MODUL_1-4.docx',
        'PANDUAN_UJIAN_CRUD_LARAVEL_MODUL_1-4.pdf',
        'serkom_guide.html'
    ]
    for ef in expected_files:
        assert os.path.exists(ef), f"File {ef} not found!"
        size = os.path.getsize(ef)
        assert size > 50000, f"File {ef} is too small: {size} bytes"
        print(f"[OK] {ef}: {size} bytes")
        
    # 2. Check Lexer Token Integrity across all code snippets
    total_tokens = 0
    total_errors = 0
    for mod_name in ['MODUL_1', 'MODUL_2', 'MODUL_3', 'MODUL_4']:
        mod = getattr(content_data, mod_name)
        for step in mod['steps']:
            fname = step['file']
            for ctype in ['default_content', 'modified_content']:
                code = step.get(ctype, '')
                if not code or code.strip() == '-':
                    continue
                lexer = build_guide.get_lexer_for_file(fname, code)
                tokens = list(lexer.get_tokens(code))
                for ttype, text in tokens:
                    total_tokens += 1
                    if ttype in Token.Error:
                        total_errors += 1
                        print(f"Lexer Error in {step['step_num']} {fname}: {repr(text)}")
                        
    assert total_errors == 0, f"Found {total_errors} lexer error tokens!"
    print(f"[OK] Lexer Token Integrity: {total_tokens} tokens parsed with 0 errors!")

    # 3. Check HTML contents
    with open('serkom_guide.html', 'r', encoding='utf-8') as f:
        raw_html = f.read()
    import html as html_lib
    html_content = html_lib.unescape(raw_html)
        
    must_have_strings = [
        'PANDUAN PRAKTIKUM & UJIAN SERKOM LARAVEL',
        'MODUL 1: SETUP DATABASE, MIGRATION, MODEL & SEEDER',
        'MODUL 2: BREEZE AUTH & CRUD MASTER MAKANAN (SISI ADMIN)',
        'MODUL 3: SISI CUSTOMER (KATALOG MENU & TRANSAKSI CHECKOUT)',
        'MODUL 4: SISI ADMIN (REKAP TRANSAKSI, UPDATE STATUS & ROUTING)',
        'TIPS & TRIK HAFALAN & STRATEGI UJIAN SERKOM',
        'SESI GRILL-ME: 10 PERTANYAAN WAWANCARA ASESOR BNSP BESERTA JAWABAN JITU',
        'KAMUS ERROR & TROUBLESHOOTING INSTAN',
        'DB::beginTransaction()',
        'DB::rollBack()',
        '@if($errors->any())',
        "old('category')",
        'confirmCheckout',
        'filterMenu',
        'php artisan storage:link',
        'Vite manifest not found',
        'SQLSTATE[42S02]',
        "protected $table = 'foods'",
        'x-app-layout',
        'bg-white',
        'border-gray-300',
        'KAMUS WAJIB: KAPAN MEMAKAI Food, food, ATAU foods'
    ]
    for s in must_have_strings:
        assert s in html_content, f"Missing string in HTML: {s}"
        print(f"[OK] Found in HTML: {s}")
        
    # 3b. Check HTML Dark Theme and Syntax Highlighting Classes
    assert '#1e1e1e' in raw_html, "Missing #1e1e1e dark background in HTML"
    assert '#18181b' in raw_html, "Missing #18181b dark container background in HTML"
    assert '-webkit-print-color-adjust: exact' in raw_html, "Missing print-color-adjust in HTML"
    syntax_classes = [
        'tok-kw', 'tok-ctrl', 'tok-class', 'tok-fn', 'tok-var',
        'tok-str', 'tok-num', 'tok-cm', 'tok-tag', 'tok-attr', 'tok-op', 'tok-punc',
        'marker-tambahkan', 'marker-ubah', 'marker-wajib', 'marker-sesuaikan', 'marker-kunci'
    ]
    for sc in syntax_classes:
        assert sc in raw_html, f"Missing syntax highlighting class {sc} in HTML"
    
    # 3c. Check In-Tag Blade syntax highlighting in HTML
    assert 'tok-fn">old<' in raw_html, "Missing in-tag old() function highlighting in HTML"
    assert 'tok-fn">route<' in raw_html, "Missing in-tag route() function highlighting in HTML"
    assert 'tok-var">$food<' in raw_html, "Missing in-tag $food variable highlighting in HTML"
    assert 'tok-var">$order<' in raw_html, "Missing in-tag $order variable highlighting in HTML"
    print("[OK] HTML Dark Theme & Advanced Syntax Highlighting verified!")

    # 4. Check DOCX structure
    doc = docx.Document('PANDUAN_SERKOM_LARAVEL_SUPER_SIMPEL.docx')
    print(f"[OK] DOCX Paragraphs: {len(doc.paragraphs)}, Tables: {len(doc.tables)}")
    assert len(doc.paragraphs) > 100, "Too few paragraphs in DOCX"
    assert len(doc.tables) >= 80, f"Expected at least 80 tables/code boxes in DOCX, got {len(doc.tables)}"
    
    # 4b. Check DOCX Dark Theme Shading, Run Colors, and Widths
    dark_cells = 0
    correct_widths = 0
    total_runs = 0
    docx_colors = set()
    for tbl in doc.tables:
        if len(tbl.rows) == 1 and len(tbl.columns) == 1:
            cell = tbl.cell(0, 0)
            tcPr = cell._tc.get_or_add_tcPr()
            shd = tcPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}shd')
            if shd is not None and shd.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill') == '1E1E1E':
                dark_cells += 1
                if abs(cell.width - Inches(6.77)) < 1000:
                    correct_widths += 1
                for p in cell.paragraphs:
                    for r in p.runs:
                        total_runs += 1
                        if r.font.color and r.font.color.rgb:
                            docx_colors.add(str(r.font.color.rgb))
                            
    assert dark_cells >= 40, f"Expected at least 40 dark code cells, found {dark_cells}"
    assert correct_widths == dark_cells, f"Expected all {dark_cells} dark cells to have 6.77in width, got {correct_widths}"
    assert total_runs >= 3000, f"Expected at least 3000 colored runs, got {total_runs}"
    assert len(docx_colors) >= 12, f"Expected at least 12 distinct syntax colors in DOCX, found {len(docx_colors)}"
    print(f"[OK] DOCX Dark Theme ({dark_cells} dark cells, width 6.77in) & Syntax Highlighting ({total_runs} runs, {len(docx_colors)} distinct colors) verified!")

    # 5. Check PDF Integrity and Pagination
    pdf_reader = PyPDF2.PdfReader('PANDUAN_SERKOM_LARAVEL_SUPER_SIMPEL.pdf')
    total_pdf_pages = len(pdf_reader.pages)
    print(f"PDF Pages: {total_pdf_pages}")
    assert total_pdf_pages == 33, f"Expected clean 33 pages (no 1-line orphan on page 34), got {total_pdf_pages}"
    
    # Check last page has meaningful content + footer note
    last_page_text = pdf_reader.pages[-1].extract_text()
    assert len(last_page_text.splitlines()) >= 10, f"Last page appears orphaned (too few lines): {len(last_page_text.splitlines())}"
    assert "Panduan Praktikum & Ujian Serkom CRUD Laravel" in last_page_text, "Missing footer note on last PDF page"
    pdf_end_norm = re.sub(r'\s+', '', ''.join(p.extract_text() for p in pdf_reader.pages[-3:]))
    assert "KAMUSERROR" in pdf_end_norm, "Missing Kamus Error near end of PDF"
    
    # Check key functional phrases across PDF
    pdf_text = ''.join(p.extract_text() for p in pdf_reader.pages)
    for kp in ['DB::beginTransaction()', 'DB::rollBack()', 'FoodController', 'OrderController', '[TAMBAHKAN]', '[UBAH]', '[WAJIB]']:
        assert kp in pdf_text, f"Missing key phrase in PDF: {kp}"
    print(f"[OK] PDF Pages ({total_pdf_pages} pages, balanced pagination) & Integrity verified!")

    print("=== ALL VERIFICATION TESTS PASSED SUCCESSFULLY! ===")

if __name__ == '__main__':
    test_verification()
