# -*- coding: utf-8 -*-
"""
Build Script for Serkom Laravel Guide
Generates both HTML/PDF and DOCX formats.
Updated: Dark Theme (#1e1e1e ala VS Code), Full Contrast Syntax Highlighting,
Semua baris kode utuh tanpa terputus, dan penanda modifikasi [TAMBAHKAN]/[UBAH] kontras.
"""

import os
import sys
import shutil
import html
import re
import subprocess
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from pygments.lexer import RegexLexer, bygroups, using
from pygments.token import (
    Token, Text, Comment, Operator, Keyword, Name, String, Number, Punctuation
)
from pygments.lexers import PhpLexer, JavascriptLexer

from content_data import (
    GUIDE_TITLE,
    GUIDE_SUBTITLE,
    TIMELINE_DATA,
    CSS_PHILOSOPHY,
    MODUL_1,
    MODUL_2,
    MODUL_3,
    MODUL_4,
    TIPS_TRICKS,
    GRILL_ME_DATA,
    KAMUS_ERROR
)

BASE_DIR = r"c:\Users\Default User\Desktop\modul serkom"
CHROME_PATH = r"C:\Program Files\Google\Chrome\Application\chrome.exe"

# ==============================================================================
# SYNTAX HIGHLIGHTING & LEXERS (DARK THEME ALA VS CODE)
# ==============================================================================

class BladeHtmlLexer(RegexLexer):
    """Robust lexer for Laravel Blade templates mixing HTML, Blade directives, and JS."""
    name = 'Blade'
    aliases = ['blade']
    filenames = ['*.blade.php']

    tokens = {
        'root': [
            (r'<!DOCTYPE.*?>', Name.Tag),
            (r'<!doctype.*?>', Name.Tag),
            (r'\{\{--', Comment.Multiline, 'blade-comment'),
            (r'\{\{', Punctuation, 'blade-echo'),
            (r'\{!!', Punctuation, 'blade-raw-echo'),
            (r'<!--[\s\S]*?-->', Comment),
            (r'(@[a-zA-Z_]\w*)(\s*)(\()', bygroups(Keyword, Text, Punctuation), 'blade-args'),
            (r'@[a-zA-Z_]\w*', Keyword),
            (r'<script.*?>', Name.Tag, 'script'),
            (r'<style.*?>', Name.Tag, 'style'),
            (r'</?[a-zA-Z0-9_\-:]+', Name.Tag, 'tag'),
            (r'&[a-zA-Z0-9#]+;', Name.Entity),
            (r'[^<@\{&]+', Text),
            (r'.', Text),
        ],
        'blade-comment': [
            (r'[^\-]+', Comment.Multiline),
            (r'--\}\}', Comment.Multiline, '#pop'),
            (r'-', Comment.Multiline),
        ],
        'blade-echo': [
            (r'\s+', Text),
            (r'\}\}', Punctuation, '#pop'),
            (r'\{\{--', Comment.Multiline, 'blade-comment'),
            (r'\b(old|route|session|asset|number_format|strtoupper|strtolower|compact|now|auth|url|back|csrf_token|__)\b', Name.Function),
            (r'\$[a-zA-Z_]\w*', Name.Variable),
            (r'->', Operator),
            (r'::', Operator),
            (r'\?\?', Operator),
            (r'==|!=|<=|>=|===|!==|<|>', Operator),
            (r'\?|:', Operator),
            (r'\+|\-|\*|\/|\%', Operator),
            (r'\.', Operator),
            (r'\'[^\']*\'', String.Single),
            (r'"[^"]*"', String.Double),
            (r'\b\d+(\.\d+)?\b', Number),
            (r'\b(true|false|null)\b', Keyword.Constant),
            (r'[a-zA-Z_]\w*', Name),
            (r'[,()\[\]]', Punctuation),
            (r'.', Text),
        ],
        'blade-raw-echo': [
            (r'\s+', Text),
            (r'!!\}', Punctuation, '#pop'),
            (r'\b(old|route|session|asset|number_format|strtoupper|strtolower|compact|now|auth|url|back|csrf_token|__)\b', Name.Function),
            (r'\$[a-zA-Z_]\w*', Name.Variable),
            (r'->', Operator),
            (r'::', Operator),
            (r'\?\?', Operator),
            (r'==|!=|<=|>=|===|!==|<|>', Operator),
            (r'\?|:', Operator),
            (r'\+|\-|\*|\/|\%', Operator),
            (r'\.', Operator),
            (r'\'[^\']*\'', String.Single),
            (r'"[^"]*"', String.Double),
            (r'\b\d+(\.\d+)?\b', Number),
            (r'\b(true|false|null)\b', Keyword.Constant),
            (r'[a-zA-Z_]\w*', Name),
            (r'[,()\[\]]', Punctuation),
            (r'.', Text),
        ],
        'blade-args': [
            (r'\(', Punctuation, '#push'),
            (r'\)', Punctuation, '#pop'),
            (r'\s+', Text),
            (r'\bas\b', Keyword),
            (r'\b(old|route|session|asset|number_format|strtoupper|compact|all|any|has|get|count|empty|selected|checked)\b', Name.Function),
            (r'\$[a-zA-Z_]\w*', Name.Variable),
            (r'->', Operator),
            (r'::', Operator),
            (r'==|!=|<=|>=|===|!==|<|>', Operator),
            (r'\?|:', Operator),
            (r'\.', Operator),
            (r'\'[^\']*\'', String.Single),
            (r'"[^"]*"', String.Double),
            (r'\b\d+\b', Number),
            (r'\b(true|false|null)\b', Keyword.Constant),
            (r'[a-zA-Z_]\w*', Name),
            (r'[,\[\]]', Punctuation),
            (r'.', Text),
        ],
        'tag': [
            (r'\s+', Text),
            (r'\{\{--', Comment.Multiline, 'blade-comment'),
            (r'\{\{', Punctuation, 'blade-echo'),
            (r'\{!!', Punctuation, 'blade-raw-echo'),
            (r'(@[a-zA-Z_]\w*)(\s*)(\()', bygroups(Keyword, Text, Punctuation), 'blade-args'),
            (r'@[a-zA-Z_]\w*', Keyword),
            (r'//.*$', Comment.Single),
            (r'[a-zA-Z0-9_\-:]+', Name.Attribute),
            (r'=', Operator),
            (r'"', String.Double, 'attr-double-quote'),
            (r'\'', String.Single, 'attr-single-quote'),
            (r'/?>', Name.Tag, '#pop'),
            (r'.', Text),
        ],
        'attr-double-quote': [
            (r'"', String.Double, '#pop'),
            (r'\{\{--', Comment.Multiline, 'blade-comment'),
            (r'\{\{', Punctuation, 'blade-echo'),
            (r'\{!!', Punctuation, 'blade-raw-echo'),
            (r'[^"\{]+', String.Double),
            (r'\{', String.Double),
        ],
        'attr-single-quote': [
            (r'\'', String.Single, '#pop'),
            (r'\{\{--', Comment.Multiline, 'blade-comment'),
            (r'\{\{', Punctuation, 'blade-echo'),
            (r'\{!!', Punctuation, 'blade-raw-echo'),
            (r'[^\'\{]+', String.Single),
            (r'\{', String.Single),
        ],
        'script': [
            (r'</script>', Name.Tag, '#pop'),
            (r'[^<]+', using(JavascriptLexer)),
            (r'<', Operator),
        ],
        'style': [
            (r'</style>', Name.Tag, '#pop'),
            (r'[^<]+', Text),
            (r'<', Text),
        ]
    }

class EnvLexer(RegexLexer):
    name = 'Env'
    tokens = {
        'root': [
            (r'//.*$', Comment.Single),
            (r'#.*$', Comment.Single),
            (r'^[A-Z0-9_]+', Name.Variable),
            (r'=', Operator),
            (r'[^#/\n\s]+', String.Single),
            (r'\s+', Text),
            (r'.', Text),
        ]
    }

class TerminalLexer(RegexLexer):
    name = 'Terminal'
    tokens = {
        'root': [
            (r'//.*$', Comment.Single),
            (r'#.*$', Comment.Single),
            (r'\b(php|composer|npm|git)\b', Keyword),
            (r'\b(artisan|require|install|run|make:model|make:seeder|migrate:fresh|storage:link|breeze:install|build)\b', Name.Function),
            (r'(-mcr|-m|-c|-r|--dev|--seed|blade|--timeout=\d+)\b', String.Single),
            (r'\b(Food|Order|OrderDetail|FoodSeeder|DatabaseSeeder)\b', Name.Class),
            (r'&&|\|', Operator),
            (r'\s+', Text),
            (r'.', Text),
        ]
    }

def get_lexer_for_file(fname, code):
    fname_lower = fname.lower()
    if '.blade.php' in fname_lower or 'blade' in fname_lower:
        return BladeHtmlLexer()
    elif '.env' in fname_lower:
        return EnvLexer()
    elif 'terminal' in fname_lower or 'command' in fname_lower:
        return TerminalLexer()
    elif '.php' in fname_lower:
        if '<?php' in code:
            return PhpLexer()
        else:
            return PhpLexer(startinline=True)
    else:
        return PhpLexer(startinline=True)

PHP_TYPES = {
    'void', 'int', 'string', 'bool', 'float', 'array', 'callable', 'iterable', 'object', 'mixed', 'never'
}

PHP_FUNCTIONS = {
    'compact', 'now', 'view', 'redirect', 'route', 'session', 'asset', 'number_format',
    'strtoupper', 'strtolower', 'old', 'back', 'response', 'abort', 'csrf_token',
    'array_filter', 'array_map', 'array_merge', 'in_array', 'count', 'empty', 'isset',
    'latest', 'paginate', 'all', 'get', 'first', 'find', 'findOrFail', 'where',
    'create', 'update', 'delete', 'destroy', 'save', 'validate', 'hasFile', 'file',
    'store', 'exists', 'beginTransaction', 'commit', 'rollBack', 'transaction',
    'hasMany', 'belongsTo', 'belongsToMany', 'hasOne', '__'
}

JS_FUNCTIONS = {
    'filterMenu', 'confirmCheckout', 'filterCategory', 'showConfirmationModal',
    'closeConfirmationModal', 'querySelectorAll', 'getElementById', 'alert',
    'parseInt', 'parseFloat', 'forEach', 'toLocaleString', 'trim', 'getAttribute',
    'addEventListener', 'preventDefault', 'appendChild', 'remove', 'add'
}

CONTROL_WORDS = {
    'return', 'if', 'else', 'elseif', 'foreach', 'endforeach', 'for', 'endfor',
    'while', 'endwhile', 'switch', 'case', 'break', 'continue', 'default',
    'try', 'catch', 'finally', 'throw', 'yield',
    '@if', '@else', '@elseif', '@endif', '@foreach', '@endforeach',
    '@forelse', '@empty', '@endforelse', '@csrf', '@method', '@include',
    '@extends', '@section', '@endsection', '@push', '@endpush'
}

MARKERS = {
    '[TAMBAHKAN]': {
        'html_class': 'marker-tambahkan',
        'docx_color': RGBColor(74, 222, 128),  # Lime Green
    },
    '[UBAH]': {
        'html_class': 'marker-ubah',
        'docx_color': RGBColor(56, 189, 248),  # Sky Blue
    },
    '[WAJIB]': {
        'html_class': 'marker-wajib',
        'docx_color': RGBColor(248, 113, 113), # Coral Red
    },
    '[SESUAIKAN]': {
        'html_class': 'marker-sesuaikan',
        'docx_color': RGBColor(192, 132, 252), # Light Purple
    },
    '[KUNCI]': {
        'html_class': 'marker-kunci',
        'docx_color': RGBColor(250, 204, 21),  # Gold
    },
}

def classify_token(ttype, text, is_html_file=False):
    stripped = text.strip()
    
    # 1. Keywords & Control Flow
    if ttype in Token.Keyword:
        if stripped in CONTROL_WORDS or stripped.startswith('@'):
            return 'control'
        if stripped in ('true', 'false', 'null', 'fn'):
            return 'keyword'
        return 'keyword'
        
    if ttype in (Token.Keyword.Constant, Token.Name.Constant):
        return 'keyword'
        
    # 2. Classes & Types
    if ttype in (Token.Name.Class, Token.Keyword.Type):
        return 'class'
        
    if stripped in PHP_TYPES:
        return 'class'
        
    if stripped.startswith('\\') and len(stripped) > 1 and stripped[1].isupper():
        return 'class'
        
    # 3. Builtins & Functions
    if ttype in (Token.Name.Function,):
        return 'function'
        
    if stripped in PHP_FUNCTIONS or stripped in JS_FUNCTIONS:
        return 'function'
        
    if ttype in (Token.Name.Builtin,):
        if stripped in ('document', 'window', 'console'):
            return 'class'
        return 'function'
        
    # 4. Variables & Attributes
    if ttype in (Token.Name.Variable, Token.Name.Variable.Instance, Token.Name.Variable.Class, Token.Name.Property):
        return 'variable'
        
    if ttype in (Token.Name.Attribute,):
        if is_html_file:
            return 'attribute'
        return 'variable'
        
    # 5. Names (Capitalized = Class, Model Properties = Variable)
    if ttype in (Token.Name.Other, Token.Name):
        if stripped in PHP_TYPES:
            return 'class'
        if stripped in PHP_FUNCTIONS or stripped in JS_FUNCTIONS:
            return 'function'
        if stripped and stripped[0].isupper():
            return 'class'
        if stripped in ('name', 'category', 'price', 'description', 'image', 'status', 'customer_name', 'table_number', 'total_price', 'quantity', 'subtotal', 'food_id', 'order_id', 'created_at', 'updated_at', 'id'):
            return 'variable'
        return 'default'
        
    # 6. Strings, Numbers, Comments, Tags, Operators, Punctuation
    if ttype in Token.Literal.String:
        return 'string'
        
    if ttype in Token.Literal.Number:
        return 'number'
        
    if ttype in Token.Comment:
        return 'comment'
        
    if ttype in (Token.Name.Tag, Token.Name.Entity):
        return 'tag'
        
    if ttype in (Token.Operator,):
        return 'operator'
        
    if ttype in (Token.Punctuation,):
        return 'punctuation'
        
    return 'default'

def highlight_code_html(code_text, file_name):
    """Tokenize and generate colored HTML spans with VS Code Dark+ color palette."""
    is_html_file = '.blade.php' in file_name.lower() or 'blade' in file_name.lower()
    lexer = get_lexer_for_file(file_name, code_text)
    tokens = list(lexer.get_tokens(code_text))
    
    out = []
    marker_pattern = '(' + '|'.join(re.escape(m) for m in MARKERS.keys()) + ')'
    
    for ttype, text in tokens:
        escaped = html.escape(text)
        
        # Check if comment has modification markers like [TAMBAHKAN]
        if ttype in Token.Comment and any(m in text for m in MARKERS):
            parts = re.split(marker_pattern, text)
            for part in parts:
                if not part:
                    continue
                if part in MARKERS:
                    out.append(f'<span class="{MARKERS[part]["html_class"]}">{html.escape(part)}</span>')
                else:
                    out.append(f'<span class="tok-cm">{html.escape(part)}</span>')
            continue
            
        category = classify_token(ttype, text, is_html_file=is_html_file)
        
        class_map = {
            'keyword': 'tok-kw',
            'control': 'tok-ctrl',
            'class': 'tok-class',
            'function': 'tok-fn',
            'variable': 'tok-var',
            'string': 'tok-str',
            'number': 'tok-num',
            'comment': 'tok-cm',
            'tag': 'tok-tag',
            'attribute': 'tok-attr',
            'operator': 'tok-op',
            'punctuation': 'tok-punc',
        }
        
        cls = class_map.get(category)
        if cls:
            out.append(f'<span class="{cls}">{escaped}</span>')
        else:
            out.append(escaped)
            
    return "".join(out)

# ==============================================================================
# 1. HTML GENERATOR (DARK CODE THEME ALA VS CODE)
# ==============================================================================
def generate_html(output_path):
    print("Generating HTML with Dark Theme & Syntax Highlighting...")
    
    html_lines = []
    html_lines.append("""<!DOCTYPE html>
<html lang="id">
<head>
    <meta charset="UTF-8">
    <title>PANDUAN PRAKTIKUM & UJIAN SERKOM LARAVEL</title>
    <style>
        * {
            -webkit-print-color-adjust: exact !important;
            print-color-adjust: exact !important;
        }
        @page {
            size: A4;
            margin: 12mm 12mm 12mm 12mm;
        }
        body {
            font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, "Helvetica Neue", Arial, sans-serif;
            color: #111;
            line-height: 1.5;
            font-size: 12px;
            margin: 0 auto;
            padding: 20px;
            max-width: 900px;
            background: #fff;
        }
        .header-box {
            text-align: center;
            border-bottom: 3px double #000;
            padding-bottom: 14px;
            margin-bottom: 22px;
        }
        .header-box h1 {
            font-size: 21px;
            margin: 0 0 6px 0;
            color: #000;
            letter-spacing: 0.5px;
        }
        .header-box .subtitle {
            font-size: 12.5px;
            color: #333;
            font-weight: 600;
        }
        .badge-target {
            display: inline-block;
            background: #f0f0f0;
            color: #000;
            border: 1px solid #000;
            padding: 4px 12px;
            border-radius: 3px;
            font-weight: bold;
            font-size: 11px;
            margin-top: 10px;
        }
        h2.section-title {
            font-size: 14.5px;
            background: #eaeaea;
            border-left: 6px solid #000;
            padding: 8px 12px;
            margin-top: 28px;
            margin-bottom: 14px;
            page-break-after: avoid;
            break-after: avoid;
            color: #000;
        }
        h3.step-title {
            font-size: 13px;
            color: #000;
            margin-top: 20px;
            margin-bottom: 6px;
            page-break-after: avoid;
            break-after: avoid;
            border-bottom: 1px dashed #666;
            padding-bottom: 4px;
        }
        p {
            margin: 6px 0;
        }
        .table-custom {
            width: 100%;
            border-collapse: collapse;
            margin: 10px 0 16px 0;
            page-break-inside: auto;
            break-inside: auto;
            font-size: 11.5px;
        }
        .table-custom tr {
            page-break-inside: avoid;
            break-inside: avoid;
        }
        .table-custom th, .table-custom td {
            border: 1px solid #333;
            padding: 5.5px 8px;
            vertical-align: top;
        }
        .table-custom th {
            background: #e2e2e2;
            font-weight: bold;
            color: #000;
        }

        /* Dark Theme Code Box (VS Code Style #1e1e1e) */
        .diff-box {
            margin: 12px 0 16px 0;
            border: 1px solid #333333;
            border-radius: 6px;
            background: #18181b;
            overflow: hidden;
            page-break-inside: auto;
            break-inside: auto;
            box-shadow: 0 4px 6px -1px rgba(0, 0, 0, 0.2);
        }
        .diff-header {
            padding: 7px 12px;
            font-size: 11px;
            font-weight: 600;
            display: flex;
            justify-content: space-between;
            align-items: center;
            page-break-after: avoid;
            break-after: avoid;
            font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif;
        }
        .diff-header-default {
            background: #27272a;
            border-bottom: 1px solid #3f3f46;
            color: #a1a1aa;
        }
        .diff-header-added {
            background: #182234;
            border-bottom: 1px solid #1e3a5f;
            color: #38bdf8;
        }
        .tag-default {
            background: #3f3f46;
            color: #e4e4e7;
            font-size: 10px;
            font-weight: 600;
            padding: 2px 8px;
            border-radius: 3px;
            letter-spacing: 0.3px;
        }
        .tag-modified {
            background: #0284c7;
            color: #ffffff;
            font-size: 10px;
            font-weight: bold;
            padding: 2px 8px;
            border-radius: 3px;
            letter-spacing: 0.3px;
        }
        .diff-body {
            padding: 0;
            background: #1e1e1e;
        }
        pre {
            background: #1e1e1e;
            color: #d4d4d4;
            padding: 10px 14px;
            font-family: "Consolas", "Courier New", monospace;
            font-size: 11px;
            line-height: 1.48;
            white-space: pre-wrap;
            word-break: break-word;
            margin: 0;
            border: none;
            border-radius: 0;
            tab-size: 4;
            page-break-inside: auto;
            break-inside: auto;
        }
        .code-added {
            border-left: 4px solid #38bdf8;
        }
        .code-default {
            border-left: 4px solid #71717a;
        }

        /* VS Code Dark+ Syntax Highlighting Colors */
        .tok-kw { color: #569cd6; font-weight: bold; }       /* Keywords: public, function, class, etc. */
        .tok-ctrl { color: #c586c0; font-weight: bold; }     /* Control flow: return, if, foreach, @if, @csrf */
        .tok-class { color: #4ec9b0; font-weight: bold; }    /* Classes: Food, Order, Request, Blueprint */
        .tok-fn { color: #dcdcaa; }                          /* Functions/Methods: store, create, validate */
        .tok-var { color: #9cdcfe; }                         /* Variables: $request, $food, $table */
        .tok-str { color: #ce9178; }                         /* Strings: 'foods', "pesanmakan" */
        .tok-num { color: #b5cea8; }                         /* Numbers: 0, 10, 25000 */
        .tok-cm { color: #6a9955; font-style: italic; }      /* Comments: // comment */
        .tok-tag { color: #569cd6; font-weight: bold; }      /* HTML tags: <table>, <form>, <x-app-layout> */
        .tok-attr { color: #9cdcfe; }                        /* HTML attributes: style, name, action */
        .tok-op { color: #d4d4d4; }                          /* Operators: ->, ::, =, + */
        .tok-punc { color: #cccccc; }                        /* Punctuation: ;, (, ), {, } */

        /* Glowing Modification Markers (Anti-Pusing & Jelas) */
        .marker-tambahkan { background: #14532d; color: #4ade80; font-weight: bold; padding: 1px 5px; border-radius: 3px; font-style: normal; }
        .marker-ubah { background: #1e3a8a; color: #38bdf8; font-weight: bold; padding: 1px 5px; border-radius: 3px; font-style: normal; }
        .marker-wajib { background: #7f1d1d; color: #f87171; font-weight: bold; padding: 1px 5px; border-radius: 3px; font-style: normal; }
        .marker-sesuaikan { background: #581c87; color: #c084fc; font-weight: bold; padding: 1px 5px; border-radius: 3px; font-style: normal; }
        .marker-kunci { background: #713f12; color: #facc15; font-weight: bold; padding: 1px 5px; border-radius: 3px; font-style: normal; }

        .callout {
            background: #fafafa;
            border-left: 4px solid #222;
            border: 1px solid #bbb;
            border-left-width: 5px;
            padding: 8px 12px;
            margin: 8px 0;
            font-size: 11.2px;
            color: #111;
            page-break-inside: avoid;
            break-inside: avoid;
        }
        .callout-title {
            font-weight: bold;
            margin-bottom: 4px;
            color: #000;
        }
        .callout-warning {
            border-left-color: #000;
            background: #fdfdfd;
        }
        .callout-success {
            border-left-color: #222;
            background: #fafafa;
        }
        .tips-card {
            border: 1px solid #555;
            padding: 9px 12px;
            margin-bottom: 10px;
            background: #ffffff;
            border-radius: 3px;
            page-break-inside: avoid;
            break-inside: avoid;
        }
        .tips-title {
            font-weight: bold;
            font-size: 12px;
            color: #000;
            margin-bottom: 4px;
        }
        .grill-card {
            border: 1px solid #222;
            padding: 10px 12px;
            margin-bottom: 12px;
            background: #fafafa;
            border-left: 5px solid #000;
            page-break-inside: avoid;
            break-inside: avoid;
        }
        .grill-q {
            font-weight: bold;
            font-size: 12px;
            color: #000;
            margin-bottom: 4px;
        }
        .grill-a {
            font-size: 11.2px;
            line-height: 1.45;
            color: #222;
        }
        @media print {
            body { padding: 0; max-width: 100%; }
            .diff-box { page-break-inside: auto; break-inside: auto; }
            pre { page-break-inside: auto; break-inside: auto; }
            .diff-header { page-break-after: avoid; break-after: avoid; }
            h2.section-title, h3.step-title { page-break-after: avoid; break-after: avoid; }
            .callout, .tips-card, .grill-card { page-break-inside: avoid; break-inside: avoid; }
            .table-custom { page-break-inside: auto; break-inside: auto; }
            .table-custom tr { page-break-inside: avoid; break-inside: avoid; }
            .footer-note { page-break-before: avoid; break-before: avoid; }
        }
    </style>
</head>
<body>
""")

    # Header
    html_lines.append(f"""
    <div class="header-box">
        <h1>{GUIDE_TITLE}</h1>
        <div class="subtitle">{GUIDE_SUBTITLE}</div>
        <div class="badge-target">🎯 TARGET: Lulus Sertifikasi Kompetensi Pemrograman Web BNSP / LSP (21 Oktober)</div>
    </div>
    """)

    # Timeline Table
    html_lines.append("""
    <h2 class="section-title">🧭 TIMELINE &amp; STRATEGI PENGERJAAN 4 JAM UJIAN</h2>
    <table class="table-custom">
        <thead>
            <tr>
                <th width="140">Alokasi Waktu</th>
                <th width="180">Modul & Target Utama</th>
                <th>Capaian Fungsional Wajib</th>
            </tr>
        </thead>
        <tbody>
    """)
    for t_time, t_mod, t_desc in TIMELINE_DATA:
        html_lines.append(f"""
            <tr>
                <td><b>{html.escape(t_time)}</b></td>
                <td><b>{html.escape(t_mod)}</b></td>
                <td>{html.escape(t_desc)}</td>
            </tr>
        """)
    html_lines.append("""
        </tbody>
    </table>
    """)

    # CSS Philosophy / Breeze Cheat Sheet
    html_lines.append("""
    <h2 class="section-title">🎨 PANDUAN STYLING BREEZE &amp; TAILWIND ESENSIAL UNTUK UJIAN</h2>
    <div class="callout callout-info">
        <div class="callout-title">💡 Manfaatkan Template Bawaan Laravel Breeze (Anti-Pusing!)</div>
        <p>Setelah menjalankan <code>php artisan breeze:install blade</code> di Modul 2, proyek Anda sudah dilengkapi Tailwind CSS dan template bawaan. <b>Jangan mengetik styling dari nol!</b> Cukup contek class standar dari file bawaan Breeze:</p>
        <ul style="margin: 8px 0 0 18px; padding: 0;">
            <li><b>Card &amp; Container Wrapper:</b> Contek dari <code>resources/views/dashboard.blade.php</code> &rarr; <code>py-12</code>, <code>max-w-7xl mx-auto sm:px-6 lg:px-8</code>, <code>bg-white overflow-hidden shadow-sm sm:rounded-lg p-6</code>.</li>
            <li><b>Form Input &amp; Label:</b> Contek dari <code>resources/views/auth/login.blade.php</code> &rarr; <code>border-gray-300 focus:border-indigo-500 focus:ring-indigo-500 rounded-md shadow-sm w-full</code>.</li>
            <li><b>Tombol Aksi (Buttons):</b> Contek dari komponen Breeze &rarr; Tombol Utama: <code>px-4 py-2 bg-gray-800 text-white rounded-md text-sm font-semibold</code>, Tombol Hapus: <code>px-3 py-1.5 bg-red-600 text-white rounded-md text-xs font-semibold</code>, Tombol Batal: <code>px-4 py-2 bg-white border border-gray-300 text-gray-700 rounded-md text-sm font-semibold</code>.</li>
            <li><b>Tabel Rapi Breeze:</b> <code>min-w-full divide-y divide-gray-200</code>, thead <code>bg-gray-50 text-gray-500 text-xs font-medium uppercase</code>, row <code>hover:bg-gray-50</code>.</li>
        </ul>
    </div>
    """)

    # Helper function for modules
    def render_modul_html(modul):
        html_lines.append(f"""
        <h2 class="section-title">{html.escape(modul['title'])}</h2>
        <p><i>{html.escape(modul['description'])}</i></p>
        """)

        if 'kamus_food_table' in modul:
            html_lines.append("""
            <h3 class="step-title" style="color: #0284c7; border-left-color: #0284c7;">📌 KAMUS WAJIB: KAPAN MEMAKAI Food, food, ATAU foods</h3>
            <div class="callout callout-warning">
                <div class="callout-title">⚠️ Mengapa Masalah 'Food vs foods' Sering Memicu Error?</div>
                <p>Kata <i>'food'</i> dalam bahasa Inggris adalah <b>uncountable noun</b> (kata benda tak terhitung). Secara default, generator Laravel terkadang menganggap jamak dari 'food' tetap 'food', sehingga saat Anda menjalankan <code>php artisan make:model Food -m</code>, Laravel bisa menghasilkan migrasi <code>create_food_table</code> (tunggal). Hal ini memicu error fatal foreign key pada <code>order_details</code> atau error <code>Table 'pesanmakan.foods' doesn't exist</code> saat Eloquent dijalankan.</p>
                <p><b>4 Formula Emas Penamaan Laravel:</b></p>
                <ul style="margin: 6px 0 0 18px; padding: 0;">
                    <li><b>Nama Tabel MySQL:</b> Wajib jamak <code>foods</code> (berakhiran s) di migrasi: <code>Schema::create('foods', ...)</code>.</li>
                    <li><b>Model Eloquent:</b> Wajib tegaskan <code>protected $table = 'foods';</code> di <code>app/Models/Food.php</code>.</li>
                    <li><b>Kolom Foreign Key:</b> Tuliskan relasi secara eksplisit: <code>$table->foreignId('food_id')->constrained('foods')->cascadeOnDelete();</code> di tabel <code>order_details</code>.</li>
                    <li><b>Variabel Controller:</b> Gunakan tunggal <code>$food</code> untuk 1 record menu dan jamak <code>$foods</code> untuk collection menu.</li>
                </ul>
            </div>
            <table class="table-custom">
                <thead>
                    <tr>
                        <th width="140">Objek</th>
                        <th width="160">Penulisan</th>
                        <th width="240">Contoh di Kode</th>
                        <th>Penjelasan</th>
                    </tr>
                </thead>
                <tbody>
            """)
            for obj, write, ex, desc in modul['kamus_food_table']:
                clean_ex = ex.replace('\\n', '\n')
                ex_fmt = "<br>".join(f"<code>{html.escape(line)}</code>" for line in clean_ex.splitlines())
                html_lines.append(f"""
                    <tr>
                        <td><b>{html.escape(obj)}</b></td>
                        <td><code>{html.escape(write)}</code></td>
                        <td>{ex_fmt}</td>
                        <td>{html.escape(desc)}</td>
                    </tr>
                """)
            html_lines.append("""
                </tbody>
            </table>
            <br>
            """)

        for step in modul['steps']:
            html_lines.append(f"""
            <h3 class="step-title">[{step['step_num']}] {html.escape(step['file'])}</h3>
            <p>{html.escape(step['action_desc'])}</p>
            """)

            # Diff container with Dark Theme
            html_lines.append("""<div class="diff-box">""")
            if step['default_content'] and step['default_content'].strip() != "-":
                hl_default = highlight_code_html(step['default_content'], step['file'])
                html_lines.append(f"""
                <div class="diff-header diff-header-default">
                    <span>📄 Kondisi Awal Bawaan Generator Laravel</span>
                    <span class="tag-default">Default File</span>
                </div>
                <div class="diff-body">
                    <pre class="code-default"><code>{hl_default}</code></pre>
                </div>
                """)

            hl_modified = highlight_code_html(step['modified_content'], step['file'])
            html_lines.append(f"""
            <div class="diff-header diff-header-added">
                <span>✏️ KODE YANG DITAMBAHKAN / DIKETIK PESERTA</span>
                <span class="tag-modified">Wajib Ditambahkan</span>
            </div>
            <div class="diff-body">
                <pre class="code-added"><code>{hl_modified}</code></pre>
            </div>
            </div>
            """)

            # Explanation
            clean_exp = html.escape(step['explanation'].replace('\\n', '\n')).replace(chr(10), '<br>')
            html_lines.append(f"""
            <div class="callout callout-success">
                <div class="callout-title">🔍 Bedah Baris & Kenapa Ditulis Begitu?</div>
                <div>{clean_exp}</div>
            </div>
            """)

    render_modul_html(MODUL_1)
    render_modul_html(MODUL_2)
    render_modul_html(MODUL_3)
    render_modul_html(MODUL_4)

    # Tips & Tricks
    html_lines.append("""
    <h2 class="section-title">🎯 TIPS &amp; TRIK HAFALAN &amp; STRATEGI UJIAN SERKOM (21 OKTOBER)</h2>
    """)
    for title, desc in TIPS_TRICKS:
        html_lines.append(f"""
        <div class="tips-card">
            <div class="tips-title">📌 {html.escape(title)}</div>
            <div style="font-size: 11.2px; line-height: 1.45;">{html.escape(desc).replace(chr(10), '<br>')}</div>
        </div>
        """)

    # Grill-Me Section (Oral Assessment Simulator)
    html_lines.append("""
    <h2 class="section-title">🔥 SESI GRILL-ME: 10 PERTANYAAN WAWANCARA ASESOR BNSP BESERTA JAWABAN JITU</h2>
    <p><i>Asesor Sertifikasi Kompetensi (BNSP/LSP) selalu melakukan tes wawancara lisan (oral assessment) untuk memastikan Anda benar-benar paham logika di balik kode yang Anda tulis. Pelajari 10 jawaban jitu di bawah ini:</i></p>
    """)
    for q_text, a_lbl, a_text in GRILL_ME_DATA:
        html_lines.append(f"""
        <div class="grill-card">
            <div class="grill-q">❓ {html.escape(q_text)}</div>
            <div class="grill-a"><b>{html.escape(a_lbl)}</b> {html.escape(a_text)}</div>
        </div>
        """)

    # Kamus Error
    html_lines.append("""
    <h2 class="section-title">🚨 KAMUS ERROR &amp; TROUBLESHOOTING INSTAN (SOLUSI 10 DETIK)</h2>
    <table class="table-custom">
        <thead>
            <tr>
                <th width="180">Pesan Error / Gejala</th>
                <th width="210">Penyebab Utama</th>
                <th>Solusi Instan 10 Detik</th>
            </tr>
        </thead>
        <tbody>
    """)
    for err, cause, sol in KAMUS_ERROR:
        html_lines.append(f"""
            <tr>
                <td><b>{html.escape(err)}</b></td>
                <td>{html.escape(cause)}</td>
                <td><b>{html.escape(sol)}</b></td>
            </tr>
        """)
    html_lines.append("""
        </tbody>
    </table>
    """)

    html_lines.append("""
    <div class="footer-note" style="text-align: center; margin-top: 10px; border-top: 1px solid #000; padding-top: 6px; color: #444; font-size: 10px; page-break-before: avoid; break-before: avoid;">
        Panduan Praktikum & Ujian Serkom CRUD Laravel — Modul 1 s.d. 4 | Semoga Sukses Ujian pada 21 Oktober!
    </div>
</body>
</html>
""")

    with open(output_path, "w", encoding="utf-8") as f:
        f.write("\n".join(html_lines))
    print(f"HTML successfully saved to: {output_path}")

# ==============================================================================
# 2. DOCX GENERATOR (DARK THEME & RUN-BASED SYNTAX HIGHLIGHTING)
# ==============================================================================
def set_cell_shading(cell, color_hex):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_border_left(cell, color_hex="38BDF8", size=24):
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="{size}" w:space="0" w:color="{color_hex}"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    cell._tc.get_or_add_tcPr().append(borders)

def set_table_borders_none(tbl):
    tblPr = tbl._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="none"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

# Colors for DOCX Syntax Highlighting
DOCX_COLORS = {
    'keyword': (RGBColor(86, 156, 214), True, False),     # #569cd6 Blue
    'control': (RGBColor(197, 134, 192), True, False),    # #c586c0 Purple
    'class': (RGBColor(78, 201, 176), True, False),       # #4ec9b0 Teal
    'function': (RGBColor(220, 220, 170), False, False),  # #dcdcaa Yellow
    'variable': (RGBColor(156, 220, 254), False, False),  # #9cdcfe Light Blue
    'string': (RGBColor(206, 145, 120), False, False),    # #ce9178 Coral
    'number': (RGBColor(181, 206, 168), False, False),    # #b5cea8 Light Green
    'comment': (RGBColor(106, 153, 85), False, True),     # #6a9955 Forest Green
    'tag': (RGBColor(86, 156, 214), True, False),         # #569cd6 Blue
    'attribute': (RGBColor(156, 220, 254), False, False), # #9cdcfe Light Blue
    'operator': (RGBColor(212, 212, 212), False, False),  # #d4d4d4
    'punctuation': (RGBColor(204, 204, 204), False, False),# #cccccc
    'default': (RGBColor(212, 212, 212), False, False),   # #d4d4d4
}

def add_styled_code_box(doc, code_text, label="KODE:", is_added=True, file_name=""):
    """Render code block inside a dark table cell (#1e1e1e) with run-level syntax highlighting."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    set_table_borders_none(tbl)
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.77)
    
    # Dark Background #1e1e1e
    set_cell_shading(cell, "1E1E1E")
    
    # Left accent border (38BDF8 Sky Blue for added, 6B7280 Gray for default)
    border_color = "38BDF8" if is_added else "6B7280"
    set_cell_border_left(cell, color_hex=border_color, size=24)
    set_cell_margins(cell, top=110, bottom=110, left=160, right=160)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    
    # Label Run
    r_lbl = p.add_run(f"[{label}]\n")
    r_lbl.bold = True
    r_lbl.font.size = Pt(8.5)
    r_lbl.font.name = "Consolas"
    if is_added:
        r_lbl.font.color.rgb = RGBColor(56, 189, 248) # Sky Blue
    else:
        r_lbl.font.color.rgb = RGBColor(156, 163, 175) # Light Gray
        
    is_html_file = '.blade.php' in file_name.lower() or 'blade' in file_name.lower()
    lexer = get_lexer_for_file(file_name, code_text)
    tokens = list(lexer.get_tokens(code_text))
    
    marker_pattern = '(' + '|'.join(re.escape(m) for m in MARKERS.keys()) + ')'
    
    for ttype, text in tokens:
        # Check for modification markers in comments
        if ttype in Token.Comment and any(m in text for m in MARKERS):
            parts = re.split(marker_pattern, text)
            for part in parts:
                if not part:
                    continue
                r = p.add_run(part)
                r.font.name = "Consolas"
                r.font.size = Pt(8.5)
                if part in MARKERS:
                    r.font.color.rgb = MARKERS[part]['docx_color']
                    r.bold = True
                    r.italic = False
                else:
                    r.font.color.rgb = RGBColor(106, 153, 85)
                    r.italic = True
            continue
            
        category = classify_token(ttype, text, is_html_file=is_html_file)
        color, bold, italic = DOCX_COLORS.get(category, DOCX_COLORS['default'])
        
        r = p.add_run(text)
        r.font.name = "Consolas"
        r.font.size = Pt(8.5)
        r.font.color.rgb = color
        if bold:
            r.bold = True
        if italic:
            r.italic = True
    
    # Small spacing after table
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(3)

def add_styled_callout(doc, title, text, style_type="info"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    set_table_borders_none(tbl)
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.77)
    bg_color = "F3F4F6" if style_type == "info" else ("F9FAFB" if style_type == "warning" else "F8FAFC")
    set_cell_shading(cell, bg_color)
    set_cell_margins(cell, top=90, bottom=90, left=140, right=140)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    
    r_title = p.add_run(f"{title}\n")
    r_title.bold = True
    r_title.font.size = Pt(9.5)
    r_title.font.name = "Arial"
    r_title.font.color.rgb = RGBColor(0, 0, 0)
        
    clean_txt = text.replace('\\n', '\n')
    r_txt = p.add_run(clean_txt)
    r_txt.font.size = Pt(9)
    r_txt.font.name = "Arial"
    r_txt.font.color.rgb = RGBColor(30, 41, 59)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(4)

def generate_docx(output_path):
    print("Generating DOCX with Dark Theme Code Blocks & Syntax Highlighting...")
    doc = docx.Document()
    
    # Page Setup (A4, 0.75 in margins)
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        
    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(3)
    r_title = p_title.add_run(GUIDE_TITLE)
    r_title.font.name = "Arial"
    r_title.font.size = Pt(17)
    r_title.bold = True
    r_title.font.color.rgb = RGBColor(0, 0, 0)
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(10)
    r_sub = p_sub.add_run(GUIDE_SUBTITLE)
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(10)
    r_sub.font.color.rgb = RGBColor(55, 65, 81)
    
    add_styled_callout(doc, "🎯 TARGET UTAMA UJIAN (21 OKTOBER)", 
                       "Modul ini dirancang agar Anda PAHAM ALUR, HAFAL KODE KUNCI, dan MAMPU menyelesaikan seluruh tes sertifikasi kompetensi (CRUD, Auth, Relasi Eloquent, Upload File, dan Transaksi DB) dalam waktu 4 jam dengan lancar.", "info")
    
    # Section: Timeline
    p_h1 = doc.add_heading(level=1)
    r_h1 = p_h1.add_run("🧭 TIMELINE & STRATEGI PENGERJAAN 4 JAM UJIAN")
    r_h1.font.name = "Arial"
    r_h1.font.size = Pt(12.5)
    r_h1.bold = True
    r_h1.font.color.rgb = RGBColor(0, 0, 0)
    
    tbl_time = doc.add_table(rows=len(TIMELINE_DATA)+1, cols=3)
    tbl_time.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_time.autofit = False
    
    col_widths = [Inches(1.8), Inches(1.8), Inches(3.2)]
    headers = ["Alokasi Waktu", "Modul & Target", "Capaian Fungsional"]
    for c_idx, h_text in enumerate(headers):
        cell = tbl_time.cell(0, c_idx)
        cell.width = col_widths[c_idx]
        set_cell_shading(cell, "E5E7EB")
        set_cell_margins(cell, top=70, bottom=70, left=90, right=90)
        p = cell.paragraphs[0]
        r = p.add_run(h_text)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(9)
        
    for r_idx, (t_time, t_mod, t_desc) in enumerate(TIMELINE_DATA, start=1):
        row = tbl_time.rows[r_idx]
        for c_idx, val in enumerate([t_time, t_mod, t_desc]):
            cell = row.cells[c_idx]
            cell.width = col_widths[c_idx]
            set_cell_margins(cell, top=60, bottom=60, left=90, right=90)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.name = "Arial"
            r.font.size = Pt(8.5)
            if c_idx < 2:
                r.bold = True

    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(4)

    # Section: CSS Philosophy / Breeze Cheat Sheet
    add_styled_callout(doc, "🎨 PANDUAN STYLING BREEZE & TAILWIND ESENSIAL", 
                       "Manfaatkan template bawaan Laravel Breeze tanpa perlu menghafal ratusan class:\n"
                       "1. Card & Container: Contek dari dashboard.blade.php -> py-12, max-w-7xl mx-auto sm:px-6 lg:px-8, bg-white overflow-hidden shadow-sm sm:rounded-lg p-6.\n"
                       "2. Form Input & Label: Contek dari auth/login.blade.php -> border-gray-300 focus:border-indigo-500 rounded-md shadow-sm w-full.\n"
                       "3. Tombol Aksi: Primary (bg-gray-800 text-white rounded-md), Danger (bg-red-600 text-white rounded-md), Batal (bg-white border border-gray-300 text-gray-700 rounded-md).\n"
                       "4. Tabel Bersih Breeze: min-w-full divide-y divide-gray-200, thead bg-gray-50 text-gray-500 text-xs font-medium uppercase.", "info")

    # Helper for rendering modules in DOCX
    def render_modul_docx(modul):
        p_m = doc.add_heading(level=1)
        r_m = p_m.add_run(modul['title'])
        r_m.font.name = "Arial"
        r_m.font.size = Pt(12.5)
        r_m.bold = True
        r_m.font.color.rgb = RGBColor(0, 0, 0)
        
        p_desc = doc.add_paragraph()
        r_desc = p_desc.add_run(modul['description'])
        r_desc.italic = True
        r_desc.font.size = Pt(9.5)

        if 'kamus_food_table' in modul:
            add_styled_callout(doc, "📌 KAMUS WAJIB: KAPAN MEMAKAI Food, food, ATAU foods",
                               "Kata 'food' dalam bahasa Inggris adalah uncountable noun (tak terhitung). Laravel kadang membuat migrasi 'create_food_table' (tunggal). Pastikan:\n"
                               "1. Migration foods: Schema::create('foods', ...)\n"
                               "2. Model Food.php: protected $table = 'foods';\n"
                               "3. Migration order_details: $table->foreignId('food_id')->constrained('foods')->cascadeOnDelete();\n"
                               "4. Controller: $food (1 item) vs $foods (collection).", "warning")
            
            tbl_k = doc.add_table(rows=len(modul['kamus_food_table'])+1, cols=4)
            tbl_k.alignment = WD_TABLE_ALIGNMENT.CENTER
            tbl_k.autofit = False
            col_w_k = [Inches(1.4), Inches(1.5), Inches(2.0), Inches(1.87)]
            headers_k = ["Objek", "Penulisan", "Contoh di Kode", "Penjelasan"]
            for c_idx, h_text in enumerate(headers_k):
                cell = tbl_k.cell(0, c_idx)
                cell.width = col_w_k[c_idx]
                set_cell_shading(cell, "E5E7EB")
                set_cell_margins(cell, top=70, bottom=70, left=90, right=90)
                p = cell.paragraphs[0]
                r = p.add_run(h_text)
                r.bold = True
                r.font.name = "Arial"
                r.font.size = Pt(9)
                
            for r_idx, (obj, write, ex, desc) in enumerate(modul['kamus_food_table'], start=1):
                row = tbl_k.rows[r_idx]
                clean_ex = ex.replace('\\n', '\n')
                for c_idx, val in enumerate([obj, write, clean_ex, desc]):
                    cell = row.cells[c_idx]
                    cell.width = col_w_k[c_idx]
                    set_cell_margins(cell, top=60, bottom=60, left=90, right=90)
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.line_spacing = 1.05
                    r = p.add_run(val)
                    if c_idx in (1, 2):
                        r.font.name = "Consolas"
                        r.font.size = Pt(8.0)
                    else:
                        r.font.name = "Arial"
                        r.font.size = Pt(8.5)
                    if c_idx == 0:
                        r.bold = True
            
            p_space = doc.add_paragraph()
            p_space.paragraph_format.space_before = Pt(4)
        
        for step in modul['steps']:
            p_s = doc.add_heading(level=2)
            r_s = p_s.add_run(f"[{step['step_num']}] {step['file']}")
            r_s.font.name = "Arial"
            r_s.font.size = Pt(10.5)
            r_s.bold = True
            r_s.font.color.rgb = RGBColor(15, 23, 42)
            
            p_act = doc.add_paragraph()
            r_act = p_act.add_run(step['action_desc'])
            r_act.font.size = Pt(9)
            
            if step['default_content'] and step['default_content'].strip() != "-":
                add_styled_code_box(doc, step['default_content'], label="📄 FILE DEFAULT LARAVEL (KONDISI AWAL)", is_added=False, file_name=step['file'])
                
            add_styled_code_box(doc, step['modified_content'], label="✏️ KODE YANG DITAMBAHKAN / DIKETIK PESERTA", is_added=True, file_name=step['file'])
            
            add_styled_callout(doc, "🔍 Bedah Baris & Kenapa Ditulis Begitu?", step['explanation'], "info")

    render_modul_docx(MODUL_1)
    render_modul_docx(MODUL_2)
    render_modul_docx(MODUL_3)
    render_modul_docx(MODUL_4)

    # Section: Tips & Tricks
    p_tips = doc.add_heading(level=1)
    r_tips = p_tips.add_run("🎯 TIPS & TRIK HAFALAN & STRATEGI UJIAN SERKOM (21 OKTOBER)")
    r_tips.font.name = "Arial"
    r_tips.font.size = Pt(12.5)
    r_tips.bold = True
    r_tips.font.color.rgb = RGBColor(0, 0, 0)
    
    for title, desc in TIPS_TRICKS:
        add_styled_callout(doc, f"📌 {title}", desc, "info")

    # Section: Grill-Me Simulator
    p_grill = doc.add_heading(level=1)
    r_grill = p_grill.add_run("🔥 SESI GRILL-ME: 10 PERTANYAAN WAWANCARA ASESOR BNSP BESERTA JAWABAN JITU")
    r_grill.font.name = "Arial"
    r_grill.font.size = Pt(12.5)
    r_grill.bold = True
    r_grill.font.color.rgb = RGBColor(0, 0, 0)
    
    for q_text, a_lbl, a_text in GRILL_ME_DATA:
        add_styled_callout(doc, f"❓ {q_text}", f"{a_lbl} {a_text}", "warning")

    # Section: Kamus Error
    p_err = doc.add_heading(level=1)
    r_err = p_err.add_run("🚨 KAMUS ERROR & TROUBLESHOOTING INSTAN (SOLUSI 10 DETIK)")
    r_err.font.name = "Arial"
    r_err.font.size = Pt(12.5)
    r_err.bold = True
    r_err.font.color.rgb = RGBColor(0, 0, 0)
    
    tbl_err = doc.add_table(rows=len(KAMUS_ERROR)+1, cols=3)
    tbl_err.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_err.autofit = False
    
    col_w_err = [Inches(2.2), Inches(2.2), Inches(2.4)]
    err_headers = ["Pesan Error / Gejala", "Penyebab Utama", "Solusi Instan 10 Detik"]
    for c_idx, h_text in enumerate(err_headers):
        cell = tbl_err.cell(0, c_idx)
        cell.width = col_w_err[c_idx]
        set_cell_shading(cell, "E5E7EB")
        set_cell_margins(cell, top=70, bottom=70, left=90, right=90)
        p = cell.paragraphs[0]
        r = p.add_run(h_text)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0, 0, 0)
        
    for r_idx, (err, cause, sol) in enumerate(KAMUS_ERROR, start=1):
        row = tbl_err.rows[r_idx]
        for c_idx, val in enumerate([err, cause, sol]):
            cell = row.cells[c_idx]
            cell.width = col_w_err[c_idx]
            set_cell_margins(cell, top=60, bottom=60, left=90, right=90)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.name = "Arial"
            r.font.size = Pt(8.5)
            if c_idx == 0:
                r.bold = True
                r.font.color.rgb = RGBColor(0, 0, 0)
            elif c_idx == 2:
                r.bold = True

    doc.save(output_path)
    print(f"DOCX successfully saved to: {output_path}")

# ==============================================================================
# 3. PDF GENERATOR (VIA CHROME HEADLESS)
# ==============================================================================
def convert_html_to_pdf(html_path, pdf_path):
    print("Converting HTML to PDF via Chrome Headless...")
    # Convert path to proper forward-slash format for Chrome
    html_url = f"file:///{os.path.abspath(html_path).replace('\\', '/')}"
    cmd = [
        CHROME_PATH,
        "--headless",
        "--disable-gpu",
        "--no-pdf-header-footer",
        f"--print-to-pdf={pdf_path}",
        html_url
    ]
    res = subprocess.run(cmd, capture_output=True, text=True)
    if res.returncode == 0 and os.path.exists(pdf_path):
        size_kb = os.path.getsize(pdf_path) / 1024
        print(f"PDF successfully generated: {pdf_path} ({size_kb:.1f} KB)")
    else:
        print(f"Failed to generate PDF. Exit code: {res.returncode}. Error: {res.stderr}")

# ==============================================================================
# 4. MAIN PIPELINE
# ==============================================================================
def main():
    print("=== STARTING SERKOM GUIDE COMPILATION ===")
    
    html_file = os.path.join(BASE_DIR, "serkom_guide.html")
    simpel_docx = os.path.join(BASE_DIR, "PANDUAN_SERKOM_LARAVEL_SUPER_SIMPEL.docx")
    simpel_pdf = os.path.join(BASE_DIR, "PANDUAN_SERKOM_LARAVEL_SUPER_SIMPEL.pdf")
    
    orig_docx = os.path.join(BASE_DIR, "PANDUAN_UJIAN_CRUD_LARAVEL_MODUL_1-4.docx")
    orig_pdf = os.path.join(BASE_DIR, "PANDUAN_UJIAN_CRUD_LARAVEL_MODUL_1-4.pdf")
    
    # 1. HTML
    generate_html(html_file)
    
    # 2. PDF
    convert_html_to_pdf(html_file, simpel_pdf)
    
    # 3. DOCX
    generate_docx(simpel_docx)
    
    # 4. Copy to original filenames as well so both are updated!
    print("Updating original files with new enhanced versions...")
    shutil.copyfile(simpel_docx, orig_docx)
    if os.path.exists(simpel_pdf):
        shutil.copyfile(simpel_pdf, orig_pdf)
        
    print("=== COMPILATION COMPLETE! ALL FILES GENERATED SUCCESSFULLY ===")

if __name__ == "__main__":
    main()
